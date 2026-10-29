import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from inventario.assets.models import HostingService


class Command(BaseCommand):
    """Importa a aba Hospedagens de uma planilha XLSX."""

    help = "Importa domínios e hospedagens a partir de um arquivo XLSX."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", type=str)
        parser.add_argument("--aba", default="Hospedagens", help="Nome da aba a importar.")

    def handle(self, *args, **options):
        path = Path(options["arquivo"])
        if not path.exists():
            raise CommandError(f"Arquivo não encontrado: {path}")

        rows = self._read_rows(path, options["aba"])
        created = updated = skipped = 0
        for row in rows:
            domain = self._value(row, "dominio")
            provider = self._value(row, "servico")
            if not domain or not provider:
                skipped += 1
                continue

            renewal_date, renewal_note = self._parse_date(row.get("vencimento"))
            cost, cost_label = self._parse_cost(row.get("valor"))
            defaults = {
                "domain": domain,
                "site": self._value(row, "site"),
                "provider": provider,
                "management_url": self._value(row, "link"),
                "renewal_date": renewal_date,
                "renewal_note": renewal_note,
                "cost": cost,
                "cost_label": cost_label,
                "billing_cycle": self._value(row, "ciclo da assinatura"),
                "notes": self._value(row, "obs"),
            }

            # Um domínio pode ter serviços independentes em provedores diferentes.
            service = HostingService.objects.filter(
                domain__iexact=domain,
                provider__iexact=provider,
                renewal_date=renewal_date,
                renewal_note__iexact=renewal_note,
                billing_cycle__iexact=defaults["billing_cycle"],
            ).first()
            if service:
                for field, value in defaults.items():
                    setattr(service, field, value)
                service.save()
                updated += 1
            else:
                HostingService.objects.create(**defaults)
                created += 1

        self.stdout.write(self.style.SUCCESS(
            f"Importação concluída: {created} criados, {updated} atualizados, {skipped} ignorados."
        ))

    def _read_rows(self, path, sheet_name):
        if path.suffix.lower() != ".xlsx":
            raise CommandError("Formato não suportado; informe uma planilha .xlsx.")

        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise CommandError(f"A aba '{sheet_name}' não existe. Abas disponíveis: {workbook.sheetnames}")

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = None
        data = []
        for row in rows:
            normalized = [self._normalize(value) for value in row]
            if "dominio" in normalized and "servico" in normalized:
                headers = normalized
                continue
            if headers is None:
                continue
            values = list(row)
            record = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            if any(value is not None and str(value).strip() for value in record.values()):
                data.append(record)
        workbook.close()
        if headers is None:
            raise CommandError("Não encontrei a linha de cabeçalho com as colunas Domínio e Serviço.")
        return data

    @staticmethod
    def _normalize(value):
        text = "" if value is None else str(value).strip().casefold()
        return "".join(
            char for char in unicodedata.normalize("NFKD", text)
            if not unicodedata.combining(char)
        )

    @classmethod
    def _value(cls, row, header):
        value = row.get(header)
        return "" if value is None else str(value).strip()

    @staticmethod
    def _parse_date(value):
        if isinstance(value, datetime):
            return value.date(), ""
        if isinstance(value, date):
            return value, ""
        text = str(value or "").strip()
        if not text:
            return None, ""
        for date_format in ("%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, date_format).date(), ""
            except ValueError:
                continue
        return None, text

    @staticmethod
    def _parse_cost(value):
        text = str(value or "").strip()
        if not text:
            return None, ""
        numeric = "".join(char for char in text if char.isdigit() or char in ".,-")
        if not numeric or not any(char.isdigit() for char in numeric):
            return None, text
        if "," in numeric and "." in numeric:
            if numeric.rfind(",") > numeric.rfind("."):
                numeric = numeric.replace(".", "").replace(",", ".")
            else:
                numeric = numeric.replace(",", "")
        elif "," in numeric:
            numeric = numeric.replace(",", ".")
        elif numeric.count(".") > 1:
            last_dot = numeric.rfind(".")
            numeric = numeric[:last_dot].replace(".", "") + numeric[last_dot:]
        try:
            return Decimal(numeric), text
        except InvalidOperation:
            return None, text
