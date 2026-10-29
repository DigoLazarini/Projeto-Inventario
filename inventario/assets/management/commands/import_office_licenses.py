import unicodedata
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from inventario.assets.models import OfficeLicense


class Command(BaseCommand):
    """Importa somente VERSÃO, Key, EMPRESA e USUÁRIO da aba Office."""

    help = "Importa licenças Office de uma planilha XLSX usando quatro colunas."
    REQUIRED_COLUMNS = ("versao", "key", "empresa", "usuario")

    def add_arguments(self, parser):
        parser.add_argument("arquivo", type=str)
        parser.add_argument("--aba", default="Licenças Office")

    def handle(self, *args, **options):
        path = Path(options["arquivo"])
        if not path.exists():
            raise CommandError(f"Arquivo não encontrado: {path}")
        if path.suffix.lower() != ".xlsx":
            raise CommandError("Formato não suportado; informe uma planilha .xlsx.")

        rows = self._read_rows(path, options["aba"])
        created = skipped = 0
        for row in rows:
            values = {
                "version": self._value(row.get("versao")),
                "key": self._value(row.get("key")),
                "company": self._value(row.get("empresa")),
                "user": self._value(row.get("usuario")),
            }
            if not any(values.values()):
                skipped += 1
                continue
            # A mesma linha da fonte não é duplicada ao executar a importação novamente.
            if OfficeLicense.objects.filter(**values).exists():
                skipped += 1
                continue
            OfficeLicense.objects.create(**values)
            created += 1

        self.stdout.write(self.style.SUCCESS(
            f"Importação concluída: {created} criadas, {skipped} ignoradas."
        ))

    def _read_rows(self, path, sheet_name):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise CommandError(
                    f"A aba '{sheet_name}' não existe. Abas disponíveis: {workbook.sheetnames}"
                )

            sheet = workbook[sheet_name]
            header_map = None
            data = []
            for row in sheet.iter_rows(values_only=True):
                normalized = [self._normalize(value) for value in row]
                if all(column in normalized for column in self.REQUIRED_COLUMNS):
                    header_map = {
                        column: normalized.index(column)
                        for column in self.REQUIRED_COLUMNS
                    }
                    continue
                if header_map is None:
                    continue
                if all(
                    index >= len(normalized) or normalized[index] == column
                    for column, index in header_map.items()
                ):
                    continue
                record = {
                    column: row[index] if index < len(row) else None
                    for column, index in header_map.items()
                }
                if any(value is not None and str(value).strip() for value in record.values()):
                    data.append(record)

            if header_map is None:
                raise CommandError(
                    "Não encontrei os cabeçalhos VERSÃO, Key, EMPRESA e USUÁRIO."
                )
            return data
        finally:
            workbook.close()

    @staticmethod
    def _normalize(value):
        text = "" if value is None else str(value).strip().casefold()
        return "".join(
            char for char in unicodedata.normalize("NFKD", text)
            if not unicodedata.combining(char)
        )

    @staticmethod
    def _value(value):
        return "" if value is None else str(value).strip()
